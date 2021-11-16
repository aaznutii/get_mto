
from glob import glob
import pandas as pd

import re
import docx
from sep_csv import create_xlsx

# получение всех файлов определенного расширения
paths = glob(r'../*/*.docx', recursive=True)


# Функция для получения данных из .docx нужных таблиц и сохранение в .csv
def get_tables():
    """
    :return: save .csv from table by .docx
    """
    header_text = ['name',
                   'room',
                   'equipment',
                   'list_of_licensed_software']
    data = []
    for fl in paths:
        doc = docx.Document(fl)
        tables = doc.tables
        for i, table in enumerate(tables):
            if re.findall(r'(№)', table.cell(0, 0).text) and len(table.columns) == 5:
                for row_i in range(1, len(table.rows)):
                    res = [tables[i].cell(row_i, col_i).text for col_i in range(1, 5)
                           if tables[i].cell(row_i, 1).text not in 'Корпус']
                    data.append(res)
                    # build = re.findall(r'\d+\b', str(res[1]))
                    # build = re.sub(r'[^\s|\s$]', r'', str(res[1]))
                    # print(build)
                    # print('=============')
        print(f'Обработан файл: {fl}')
    df = pd.DataFrame(data=data, columns=header_text)
    df.to_csv('base_mto_su.csv', encoding='utf-8')


def main():
    get_tables()
    create_xlsx()


if __name__ == '__main__':
    main()
